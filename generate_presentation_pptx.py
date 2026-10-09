"""Generates a professional 12-slide PowerPoint presentation (.pptx) with embedded charts."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUTPUT_PPTX = "presentation.pptx"
PLOTS_DIR = "images/generated_analysis_plots"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
PRIMARY_COLOR = RGBColor(30, 58, 138)       # Deep Blue
SECONDARY_COLOR = RGBColor(16, 185, 129)   # Emerald Green
TEXT_DARK = RGBColor(31, 41, 55)           # Gray 800
BG_LIGHT = RGBColor(248, 250, 252)         # Slate 50


def add_slide_header(slide, title_text, category_text="ENGINEERING MINI-PROJECT"):
    # Category Tracker
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.4))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = SECONDARY_COLOR

    # Title
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = PRIMARY_COLOR


# Slide 1: Title Slide
blank_layout = prs.slide_layouts[6]
s1 = prs.slides.add_slide(blank_layout)

title_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.0))
tf1 = title_box.text_frame
tf1.word_wrap = True

p1 = tf1.paragraphs[0]
p1.text = "Sleep Disorder Classification Using Machine Learning"
p1.font.size = Pt(36)
p1.font.bold = True
p1.font.color.rgb = PRIMARY_COLOR
p1.alignment = PP_ALIGN.LEFT

p2 = tf1.add_paragraph()
p2.text = "A Multi-Class Supervised Learning Framework for Lifestyle & Sleep Health Assessment"
p2.font.size = Pt(20)
p2.font.color.rgb = RGBColor(75, 85, 99)
p2.space_before = Pt(16)

p3 = tf1.add_paragraph()
p3.text = "Department of Computer Science & Engineering | Engineering Mini-Project"
p3.font.size = Pt(14)
p3.font.color.rgb = SECONDARY_COLOR
p3.space_before = Pt(30)


# Helper for standard two-column content slide
def create_content_slide(title, category, bullet_points, image_path=None):
    slide = prs.slides.add_slide(blank_layout)
    add_slide_header(slide, title, category)

    content_width = Inches(6.5) if image_path and os.path.exists(image_path) else Inches(11.5)
    tbox = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), content_width, Inches(5.2))
    tf = tbox.text_frame
    tf.word_wrap = True

    for i, bp in enumerate(bullet_points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {bp}"
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(14)

    if image_path and os.path.exists(image_path):
        slide.shapes.add_picture(image_path, Inches(7.5), Inches(1.7), width=Inches(5.0))

    return slide


# Slide 2: Introduction
create_content_slide(
    "1. Introduction & Background",
    "CONTEXT & RELEVANCE",
    [
        "Sleep is an essential biological process directly affecting cognitive function, immune resilience, and cardiovascular health.",
        "Over 30% of global adults suffer from chronic sleep disorders such as Insomnia and Sleep Apnea.",
        "Traditional clinical diagnosis relies on overnight Polysomnography (PSG), which is expensive, time-intensive, and resource-constrained.",
        "Machine Learning provides non-invasive, accessible screening tools by analyzing everyday physiological and behavioral indicators."
    ]
)

# Slide 3: Problem Statement
create_content_slide(
    "2. Problem Statement",
    "CHALLENGE & MOTIVATION",
    [
        "Diagnostic Delay: Millions remain undiagnosed due to the friction of laboratory sleep tests and lack of routine sleep screenings.",
        "Multi-Factor Complexity: Sleep disorders stem from interconnected lifestyle, occupational, and physiological metrics.",
        "Clinical Triaging Need: Healthcare providers require an automated, educational multi-class classification tool to stratify risk into None, Insomnia, and Sleep Apnea.",
        "Educational Demonstration: Building a reproducible, transparent machine learning application tailored for engineering and clinical awareness."
    ]
)

# Slide 4: Objectives
create_content_slide(
    "3. Project Objectives",
    "GOALS & SCOPE",
    [
        "Multi-Class Classification: Accurately predict whether a subject has No Disorder, Insomnia, or Sleep Apnea.",
        "End-to-End Pipeline: Build a Scikit-Learn ColumnTransformer and Pipeline to prevent data leakage and guarantee training-inference parity.",
        "Comprehensive Benchmarking: Compare 5 prominent classification algorithms using 5-Fold Stratified Cross-Validation.",
        "Interactive Deployment: Develop a polished Streamlit web application providing real-time probability estimates and statistical dashboards."
    ]
)

# Slide 5: Dataset
create_content_slide(
    "4. Dataset Overview & Characteristics",
    "DATA & EXPLORATION",
    [
        "Source: Kaggle Sleep Health and Lifestyle Dataset (374 participant records, 13 features).",
        "Target Classes: None (58.6%, 219 samples), Sleep Apnea (20.9%, 78 samples), Insomnia (20.6%, 77 samples).",
        "Key Demographic & Lifestyle Attributes: Age (27-59), Gender, Occupation (11 roles), Physical Activity Level, Daily Steps.",
        "Key Sleep & Physiological Attributes: Sleep Duration, Quality of Sleep, Stress Level, BMI Category, Heart Rate, Blood Pressure."
    ],
    f"{PLOTS_DIR}/01_target_distribution.png"
)

# Slide 6: Preprocessing
create_content_slide(
    "5. Data Preprocessing & Feature Engineering",
    "METHODOLOGY & PIPELINE",
    [
        "Handling Missing Labels: Mapped 219 null entries in 'Sleep Disorder' to 'None' reflecting healthy subjects.",
        "Blood Pressure Decomposition: Converted compound string (e.g. '126/83') into continuous Systolic_BP and Diastolic_BP features.",
        "Category Normalization: Standardized synonymous BMI labels ('Normal Weight' -> 'Normal').",
        "Leakage Prevention: Fitted SimpleImputer, StandardScaler, and OneHotEncoder exclusively on the 80% training set within a Scikit-Learn Pipeline."
    ],
    f"{PLOTS_DIR}/07_bmi_category_distribution.png"
)

# Slide 7: Algorithms
create_content_slide(
    "6. Evaluated Machine Learning Algorithms",
    "MODEL SELECTION",
    [
        "1. Logistic Regression: Baseline linear model with balanced class weighting for calibrated probabilities.",
        "2. Decision Tree Classifier: Non-linear tree-based partitions with depth regularization to curb variance.",
        "3. Random Forest Classifier: Ensemble of 100 bagged trees mitigating overfitting and capturing multi-way feature interactions.",
        "4. K-Nearest Neighbors (KNN): Distance-weighted neighborhood classification over scaled feature vectors.",
        "5. Gradient Boosting: Sequentially boosted decision stumps optimizing multi-class log-loss."
    ]
)

# Slide 8: Architecture
create_content_slide(
    "7. System Architecture & Workflow",
    "SYSTEM DESIGN",
    [
        "Raw Data Ingestion: Patient data from CSV or Streamlit user input form.",
        "Custom Feature Engineering: FeatureEngineer transformer parses BP, standardizes BMI, drops Person ID.",
        "ColumnTransformer: Standardizes 9 numerical features; one-hot encodes 3 categorical features.",
        "Ensemble Classifier: Saved Random Forest model pipeline (sleep_disorder_pipeline.joblib).",
        "Inference & Presentation: Real-time class prediction, class probabilities, and visual guidance via Streamlit UI."
    ],
    f"{PLOTS_DIR}/12_correlation_heatmap.png"
)

# Slide 9: Evaluation
create_content_slide(
    "8. Model Evaluation & Comparison",
    "BENCHMARK RESULTS",
    [
        "Champion Model: Random Forest achieved highest Test Macro F1 (0.9352) and Test Accuracy (96.00%).",
        "Gradient Boosting achieved 94.67% accuracy and 0.9214 Macro F1.",
        "KNN & Logistic Regression achieved 93.33% accuracy.",
        "Prioritized Metric: Macro F1-Score was favored over accuracy to protect sensitivity on minority classes (Insomnia & Sleep Apnea)."
    ],
    f"{PLOTS_DIR}/model_comparison.png"
)

# Slide 10: Streamlit App
create_content_slide(
    "9. Streamlit Web Application Interface",
    "USER INTERFACE & DEPLOYMENT",
    [
        "5 Navigation Sections: Home, Predict, Dataset Insights, Model Performance, About Project.",
        "Interactive Predict Form: Sliders and select boxes with instant validation and compound BP synthesis.",
        "Probability Confidence: Multi-class probability breakdown explaining relative model certainty.",
        "Ethical Medical Disclaimer: Prominently displayed warning that predictions are educational and non-diagnostic."
    ]
)

# Slide 11: Results & Discussion
create_content_slide(
    "10. Clinical & Statistical Insights",
    "RESULTS & DISCUSSION",
    [
        "Perfect Healthy Classification: 100% Precision and 100% Recall on the 'None' class on test data.",
        "Sleep Apnea Sensitivity: 93.8% Recall, strongly correlated with elevated BMI and Systolic BP > 130 mmHg.",
        "Insomnia Patterns: 86.7% Recall and 92.9% Precision, strongly driven by Sleep Duration < 6.5 hrs and Stress >= 7.",
        "Confusion Matrix: Only 3 minor misclassifications out of 75 held-out test subjects."
    ],
    f"{PLOTS_DIR}/confusion_matrix.png"
)

# Slide 12: Conclusion & Future Scope
create_content_slide(
    "11. Conclusion & Future Enhancements",
    "CONCLUSION & ROADMAP",
    [
        "Conclusion: Successfully built an educational, highly accurate (96%), and reproducible sleep disorder classification system.",
        "Wearable Integration: Ingesting continuous PPG and accelerometer telemetry from smartwatches (Fitbit, Apple Watch).",
        "Deep Learning: Extending to raw 1D-CNN or LSTM architectures on polysomnography time-series data.",
        "Clinical Validation: Expanding the dataset to multi-hospital cohorts across diverse age groups and ethnicities."
    ]
)

prs.save(OUTPUT_PPTX)
print(f"Successfully generated PowerPoint presentation at '{OUTPUT_PPTX}'.")
